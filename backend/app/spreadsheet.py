"""Turn an Excel statement (.xlsx, optionally password-protected) into CSV bytes.

Banks such as KVB send statements as locked Excel files. Converting them to CSV here lets the
normal importer do the rest, so there is one parser to trust. The file and the password are
used in memory only: nothing is written to disk or stored.
"""

import csv
import io
import zipfile
from datetime import date, datetime

from app.importer import MAX_ROWS, StatementError

OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"  # Office's encrypted-file wrapper (also the old .xls format)
ZIP_MAGIC = b"PK\x03\x04"  # a plain .xlsx is a zip
MAX_UNZIPPED_BYTES = 30 * 1024 * 1024  # a 2 MB upload must not inflate into gigabytes (zip bomb)


def is_spreadsheet(content: bytes) -> bool:
    return content.startswith((OLE_MAGIC, ZIP_MAGIC))


def _decrypt(content: bytes, password: str | None) -> bytes:
    import msoffcrypto  # imported here: only needed for locked files

    try:
        office = msoffcrypto.OfficeFile(io.BytesIO(content))
        if not office.is_encrypted():
            raise StatementError(
                "That is an old-style Excel file (.xls). Open it in Excel and save it as .xlsx or CSV, then upload again."
            )
        if not password:
            raise StatementError("This Excel file is password-protected. Enter its password and upload again.")
        office.load_key(password=password)
        out = io.BytesIO()
        office.decrypt(out)
        return out.getvalue()
    except StatementError:
        raise
    except Exception:
        # msoffcrypto raises assorted errors for a wrong password or a damaged file
        raise StatementError("Could not unlock the file. Check the password and try again.")


def _cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        # midnight means the cell held only a date; otherwise keep the time (read later as IST)
        return value.strftime("%Y-%m-%d") if value.time() == datetime.min.time() else value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else repr(value)
    return str(value).strip()


def spreadsheet_to_csv(content: bytes, password: str | None = None) -> bytes:
    if content.startswith(OLE_MAGIC):
        content = _decrypt(content, password)
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(info.file_size for info in archive.infolist()) > MAX_UNZIPPED_BYTES:
                raise StatementError("That spreadsheet is too large. Export a shorter date range.")
    except zipfile.BadZipFile:
        raise StatementError("That doesn't look like a valid Excel file.")

    from openpyxl import load_workbook

    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        sheet = workbook.worksheets[0]
        out = io.StringIO()
        writer = csv.writer(out)
        for count, row in enumerate(sheet.iter_rows(values_only=True)):
            if count > MAX_ROWS + 60:
                raise StatementError(f"Too many rows (the limit is {MAX_ROWS:,}). Split the statement into smaller files.")
            writer.writerow([_cell_text(v) for v in row])
        workbook.close()
    except StatementError:
        raise
    except Exception:
        raise StatementError("The Excel file couldn't be read. Try exporting it as CSV instead.")
    return out.getvalue().encode("utf-8")
